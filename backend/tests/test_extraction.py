import io
import pytest
from app.services.extraction_service import (
    ExtractionService,
    CorruptFileError,
    EmptyDocumentError,
    ExtractionError,
)
from reportlab.pdfgen import canvas
import docx
from PIL import Image, ImageDraw


def create_sample_pdf(text_lines):
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    for page_idx, line in enumerate(text_lines):
        c.drawString(100, 750, line)
        if page_idx < len(text_lines) - 1:
            c.showPage()
    c.save()
    return buf.getvalue()


def create_sample_docx(paragraphs, table_data=None):
    doc = docx.Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    if table_data:
        table = doc.add_table(rows=len(table_data), cols=len(table_data[0]))
        for r_idx, row in enumerate(table_data):
            for c_idx, val in enumerate(row):
                table.cell(r_idx, c_idx).text = val
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def create_sample_image(text="Invoice 98765"):
    img = Image.new("RGB", (320, 100), color=(255, 255, 255))
    d = ImageDraw.Draw(img)
    d.text((20, 40), text, fill=(0, 0, 0))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def test_pdf_extraction_success():
    pdf_bytes = create_sample_pdf([
        "DocRAG is an enterprise intelligence platform.",
        "Page two covers pgvector and RAG retrieval architecture."
    ])
    pages = ExtractionService.extract(pdf_bytes, "architecture.pdf")
    assert len(pages) == 2
    assert pages[0].page_number == 1
    assert "DocRAG" in pages[0].text
    assert pages[0].metadata["extraction_method"] == "digital_pdf"
    assert pages[1].page_number == 2
    assert "pgvector" in pages[1].text


def test_pdf_corrupt_file_error():
    corrupt_bytes = b"NOT_A_PDF_CONTENT"
    with pytest.raises(CorruptFileError):
        ExtractionService.extract(corrupt_bytes, "bad.pdf")


def test_empty_document_error():
    with pytest.raises(EmptyDocumentError):
        ExtractionService.extract(b"", "empty.pdf")


def test_docx_extraction_success():
    docx_bytes = create_sample_docx(
        ["DocRAG supports enterprise DOCX documents."],
        table_data=[["Metric", "Value"], ["Precision", "99%"]]
    )
    pages = ExtractionService.extract(docx_bytes, "sample.docx")
    assert len(pages) >= 1
    assert "DocRAG supports enterprise DOCX" in pages[0].text
    assert "| Metric | Value |" in pages[0].text
    assert "| Precision | 99% |" in pages[0].text
    assert pages[0].metadata["extraction_method"] == "docx"


def test_image_ocr_extraction():
    img_bytes = create_sample_image("Receipt Total $450")
    pages = ExtractionService.extract(img_bytes, "receipt.png")
    assert len(pages) == 1
    assert "Receipt" in pages[0].text or "Total" in pages[0].text
    assert pages[0].metadata["extraction_method"] == "ocr_image"
    assert pages[0].metadata["is_scanned"] is True
    assert pages[0].metadata["ocr_confidence"] > 0.0


def test_unsupported_file_type():
    with pytest.raises(ExtractionError):
        ExtractionService.extract(b"Plain text file", "sample.txt")
